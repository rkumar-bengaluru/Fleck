"""Build parsing fixtures rate-card.xlsx and rate-card.pptx (deterministic). Run: uv run --package fleck-api python api/tests/fixtures/make_xlsx.py"""
import datetime as dt
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).parent
FIXED = dt.datetime(2026, 4, 1)
ROWS = [("Engagement manager", 9800, 22000), ("Senior engineer", 5200, 13500), ("ML engineer", 6800, 17000), ("QA engineer", 3600, 10000)]


def _rezip(path: Path) -> None:
    """Rewrite the archive with fixed entry timestamps so output bytes are stable."""
    tmp = Path(tempfile.mkdtemp()) / "x"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in sorted(zin.namelist()):
            zi = zipfile.ZipInfo(n, (2026, 4, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            data = zin.read(n)
            if n == "docProps/core.xml":  # openpyxl stamps modified=now on save
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-04-01T00:00:00Z", data)
            zout.writestr(zi, data)
    shutil.move(tmp, path)


def xlsx() -> None:
    wb = Workbook()
    wb.properties.creator, wb.properties.created, wb.properties.modified = "Northbeam Digital", FIXED, FIXED
    ws = wb.active
    ws.title = "Rate card"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Northbeam Digital: rate card FY2026 (USD per person per month)"
    ws["A1"].font, ws["A1"].alignment = Font(bold=True, size=14), Alignment(horizontal="center")
    ws.append(["Role", "Offshore (Bengaluru)", "Onsite (US/UK)"])
    for r in ROWS:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=3, min_col=2, max_col=3):
        for c in row:
            c.number_format = '"$"#,##0'
    p = OUT / "rate-card.xlsx"
    wb.save(p)
    _rezip(p)


def pptx() -> None:
    prs = Presentation()
    prs.core_properties.author, prs.core_properties.created, prs.core_properties.modified = "Northbeam Digital", FIXED, FIXED
    prs.core_properties.last_modified_by = "Northbeam Digital"
    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Rate card FY2026 (USD per person per month)"
    t = s.shapes.add_table(len(ROWS) + 1, 3, Inches(0.5), Inches(1.8), Inches(9), Inches(2.5)).table
    for j, h in enumerate(["Role", "Offshore (Bengaluru)", "Onsite (US/UK)"]):
        t.cell(0, j).text = h
    for i, r in enumerate(ROWS, 1):
        t.cell(i, 0).text, t.cell(i, 1).text, t.cell(i, 2).text = r[0], f"{r[1]:,}", f"{r[2]:,}"
    p = OUT / "rate-card.pptx"
    prs.save(p)
    _rezip(p)


if __name__ == "__main__":
    xlsx()
    pptx()
